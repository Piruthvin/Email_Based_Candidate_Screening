import io
import os
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Candidate, RankingResult, RankingRun


class ReportService:
    """
    Generates formatted Excel reports (.xlsx) from frozen ranking runs.
    """

    async def generate_ranking_excel(
        self,
        db: AsyncSession,
        run_id: str,
        columns: Optional[List[str]] = None,
    ) -> tuple[bytes, str, int]:
        import uuid
        try:
            run_uuid = uuid.UUID(str(run_id))
        except (ValueError, TypeError):
            raise ValueError(f"Invalid UUID for ranking run: {run_id}")

        q_run = select(RankingRun).where(RankingRun.id == run_uuid)
        run = (await db.execute(q_run)).scalars().first()
        if not run:
            raise ValueError(f"Ranking run {run_id} not found")

        q_res = (
            select(RankingResult, Candidate)
            .join(Candidate, Candidate.id == RankingResult.candidate_id)
            .where(RankingResult.run_id == run.id)
            .order_by(RankingResult.rank_position.asc())
        )
        results = (await db.execute(q_res)).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Ranked Shortlist"

        # Headers
        default_headers = [
            "Rank",
            "Candidate Name",
            "Final Score",
            "Details Score",
            "Resume Score",
            "Experience (Yrs)",
            "Notice Period (Days)",
            "Current Location",
            "Missing Skills",
            "Justification Reason",
        ]
        ws.append(default_headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_num in range(1, len(default_headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Rows
        for r_item, cand in results:
            row_data = [
                r_item.rank_position,
                cand.full_name,
                float(r_item.final_score),
                float(r_item.details_score),
                float(r_item.resume_score) if r_item.resume_score is not None else "N/A",
                float(cand.experience_years) if cand.experience_years else "N/A",
                cand.notice_days_max if cand.notice_days_max is not None else "N/A",
                cand.location or "N/A",
                ", ".join(r_item.missing_skills) if r_item.missing_skills else "None",
                r_item.reason or "",
            ]
            ws.append(row_data)

        # Style data rows and adjust column widths
        for row in ws.iter_rows(min_row=2, max_row=len(results) + 1):
            for cell in row:
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        file_bytes = output.getvalue()
        filename = f"ranking_report_{str(run.id)[:8]}.xlsx"

        return file_bytes, filename, len(results)


report_service = ReportService()
